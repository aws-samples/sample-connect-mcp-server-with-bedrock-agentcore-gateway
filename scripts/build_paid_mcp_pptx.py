"""Generate `docs/slides/paid-mcp/paid-mcp.pptx` — the 3-page paid-MCP deck.

A script rather than a hand-edited file for two reasons learned the hard way:

- Editing an existing deck in place (delete slide N, add a replacement) corrupts the package:
  python-pptx removes the slide id but the part stays, and saving writes a duplicate
  `ppt/slides/slideN.xml` that PowerPoint and LibreOffice both refuse to open. Rebuilding every
  slide from scratch has no such failure mode.
- Estimating text width in EMU for chip/box shapes silently overflows with CJK text at small point
  sizes — the fix is to stop drawing boxes around measured text, which only stays fixed if the
  layout lives in code someone can re-run.

Fonts are deliberately the ones every machine has (Georgia / Arial / Consolas). The HTML deck uses
Bespoke Serif + Satoshi; embedding those in a pptx would reflow on any machine that lacks them.

Run: `uv run --with python-pptx python scripts/build_paid_mcp_pptx.py`
Check: render it (`soffice --headless --convert-to pdf`) and look, rather than trusting the code.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "slides" / "paid-mcp" / "paid-mcp.pptx"

# Dark theme matching the AgentCore Payments deck: pure black stage, vivid purple emphasis,
# navy cards with purple keylines, mint for mono labels. White is the body colour, not the accent.
PURPLE = RGBColor(0x8C, 0x33, 0xFF)  # vivid purple fill
PURPLE_DIM = RGBColor(0x5B, 0x21, 0xC7)
LILAC = RGBColor(0xC9, 0xA9, 0xFF)  # purple that stays readable on black
MINT = RGBColor(0x7D, 0xF9, 0xC0)  # small mono labels / taglines
INK = RGBColor(0xFF, 0xFF, 0xFF)  # body text on black
INK2 = RGBColor(0xC7, 0xC7, 0xD8)
INK3 = RGBColor(0x8A, 0x8A, 0xA3)
LINE = RGBColor(0x2C, 0x2C, 0x40)  # hairline on black
PAPER = RGBColor(0x00, 0x00, 0x00)  # stage
NAVY = RGBColor(0x0E, 0x1A, 0x47)  # card fill
GHOST = RGBColor(0x2A, 0x1B, 0x4D)  # oversized numerals behind rows
CODEBG = RGBColor(0x0A, 0x0A, 0x14)
HLBG = RGBColor(0x2A, 0x1F, 0x4E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# code token colours
CK, CS, CN, CP, CTXT = (
    LILAC,
    MINT,
    RGBColor(0xFF, 0xCE, 0x86),
    RGBColor(0x8B, 0x84, 0xA8),
    RGBColor(0xE8, 0xE2, 0xFF),
)
MONO, DISP, BODY = "Consolas", "Georgia", "Arial"


def e(inches: float) -> Emu:
    return Emu(int(inches * 914400))


_TOKEN = re.compile(
    r'(//.*$)|("(?:[^"\\]|\\.)*"\s*:)|("(?:[^"\\]|\\.)*")'
    r"|(\btrue\b|\bfalse\b|\b\d+\b)|([{}\[\],:]+)|(\s+)|([^\s{}\[\],:\"]+)"
)


def colorize(line: str, size: float = 9) -> list[tuple[str, dict]]:
    """Split a JSON-ish line into coloured runs (keys, strings, numbers, comments)."""
    runs: list[tuple[str, dict]] = []
    for m in _TOKEN.finditer(line):
        comment, key, string, num, punct, space, other = m.groups()
        if comment:
            runs.append((comment, {"color": CP, "font": MONO, "size": size, "italic": True}))
        elif key:
            runs.append((key, {"color": CK, "font": MONO, "size": size}))
        elif string:
            runs.append((string, {"color": CS, "font": MONO, "size": size}))
        elif num:
            runs.append((num, {"color": CN, "font": MONO, "size": size}))
        elif punct:
            runs.append((punct, {"color": CP, "font": MONO, "size": size}))
        else:
            runs.append((space or other, {"color": CTXT, "font": MONO, "size": size}))
    return runs or [(line, {"color": CTXT, "font": MONO, "size": size})]


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = e(13.333)
        self.prs.slide_height = e(7.5)
        self.w, self.h = self.prs.slide_width, self.prs.slide_height
        self.s = None  # current slide

    def slide(self) -> None:
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.box(0, 0, self.w, self.h, fill=PAPER)

    def box(self, x, y, w, h, fill=None, line=None, lw=1):
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = line
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def text(self, x, y, w, h, lines, size=11, color=INK, font=BODY, space=1.3, align=None):
        tb = self.s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        first = True
        for line in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align or PP_ALIGN.LEFT
            p.line_spacing = space
            for text, st in line if isinstance(line, list) else [(line, {})]:
                r = p.add_run()
                r.text = text
                r.font.size = Pt(st.get("size", size))
                r.font.bold = st.get("bold", False)
                r.font.italic = st.get("italic", False)
                r.font.color.rgb = st.get("color", color)
                r.font.name = st.get("font", font)
        return tb

    def footer(self, left: str, page: str) -> None:
        fy = self.h - e(0.62)
        self.box(e(0.72), fy, self.w - e(1.44), Emu(1), fill=LINE)
        self.text(e(0.72), fy + e(0.12), e(8.4), e(0.25), [left], size=8.5, color=INK3, font=MONO)
        self.text(
            self.w - e(1.9),
            fy + e(0.12),
            e(1.2),
            e(0.25),
            [page],
            size=8.5,
            color=INK3,
            font=MONO,
            align=PP_ALIGN.RIGHT,
        )

    def head(self, eyebrow: str, title: str) -> None:
        self.text(
            e(0.72),
            e(0.5),
            e(9),
            e(0.3),
            [[(eyebrow, {"size": 11, "color": LILAC, "font": MONO, "bold": True})]],
        )
        self.text(
            e(0.72), e(0.78), e(9.6), e(0.8), [[(title, {"size": 30, "color": INK, "font": DISP})]]
        )


REQS: list[tuple[str, str, str, bool]] = [
    (
        "01",
        "一个 MCP 端点",
        "POST /mcp · JSON-RPC 2.0,实现 initialize / tools/list / tools/call",
        False,
    ),
    (
        "02",
        "schema 里留一个 headers",
        "Buyers 的 agent 全靠 inputSchema 决定怎么调;这个对象参数是付款证明的回传口",
        True,
    ),
    (
        "03",
        "路由前挂 x402 中间件",
        "在业务逻辑之前声明 scheme、network、payTo、金额与 asset——定价就是一个数字",
        False,
    ),
    (
        "04",
        "402 走 MCP payload",
        "challenge 放 structuredContent,再配一条 PAYMENT_REQUIRED: 文本标记兜底",
        True,
    ),
    (
        "05",
        "结算凭证也放 payload",
        "Gateway 会丢掉 responseHeaders,只写在头里 Buyers 拿不到链上凭证",
        True,
    ),
    (
        "06",
        "收款账户 + Gateway target",
        "payTo 需已有该稳定币代币账户;建 mcpServer target 让任何 agent 都能发现你",
        False,
    ),
]

SNIPPETS: list[tuple[str, str, list[str], list[int]]] = [
    (
        "02",
        "TOOLS/LIST · ONE PAID TOOL",
        [
            '{ "name": "price", "description": "USD price · 0.001 USDC",',
            '  "inputSchema": { "type": "object", "properties": {',
            '      "coin":    { "type": "string" },',
            '      "headers": { "type": "object" }   // 付款证明回传口',
            "  } } }",
        ],
        [3],
    ),
    (
        "04",
        "UNPAID → PAYMENT REQUIRED",
        [
            '{ "content": [{ "text": "PAYMENT_REQUIRED: {…}" }],',
            '  "structuredContent": { "x402Version": 2,',
            '    "accepts": [{ "scheme": "exact", "amount": "1000",',
            '      "network": "solana:EtWT…", "payTo": "6m9U…",',
            '      "asset": "4zMM…USDC" }] } }',
        ],
        [1],
    ),
    (
        "05",
        "PAID → RESULT + RECEIPT",
        [
            '{ "structuredContent": { "coin": "bitcoin", "usd": 77515,',
            '    "_x402": { "success": true, "transaction": "3Q3XLo…" }',
            "  } }",
        ],
        [1],
    ),
]

ASSUMPTIONS: list[tuple[str, str, str, str, str]] = [
    (
        "01",
        "调用结果",
        "一次 tools/call 只有成功或报错两种归宿,客户端照着这两种写分支就够了。",
        "多出第三种:要求付款。challenge 在 structuredContent 里,"
        + "客户端据此去签名,而不是当成错误上报。",
        "OK / ERROR / PAY-FIRST",
    ),
    (
        "02",
        "凭证在哪",
        "认证在传输层:Authorization: Bearer … 或 OAuth,一次授权长期有效。",
        "在工具参数里:headers 参数带一张一次性支付证明。Gateway 转发参数、不能注入 HTTP 头——"
        + "这是它能穿过 Gateway 的唯一办法。",
        "PER-CALL, NOT PER-SESSION",
    ),
    (
        "03",
        "你是谁",
        "先注册 client、发 token,再维护配额表与账单关系,服务端必须认识每个调用方。",
        "没有账号、没有 key、没有配额表。授权就是链上一次 USDC 转账,结算成功才执行——"
        + "服务端不需要认识调用方。",
        "NO SIGNUP, NO QUOTA TABLE",
    ),
    (
        "04",
        "重试语义",
        "同一个请求可以原样重放,幂等,失败了重试就好。",
        "证明约 60 秒有效,必须用新连接重试,且每次重试都要一张新证明——过期证明等同未付款。",
        "PROOF EXPIRES, REPLAY DOES NOT WORK",
    ),
]

WALLETS: list[tuple[str, str, list[list[tuple[str, dict]]], str]] = [
    (
        "Coinbase CDP",
        "instrument network: ETHEREUM · SOLANA",
        [
            [("eip155:1 Ethereum  ·  eip155:8453 Base", {})],
            [("base-sepolia", {"color": MINT}), ("  ← 测试网首选", {"color": INK3})],
            [("Solana mainnet", {})],
        ],
        "凭证:API Key ID + Secret + Wallet Secret。",
    ),
    (
        "Stripe(Privy)",
        "instrument network: ETHEREUM · SOLANA",
        [
            [("eip155:1 Ethereum  ·  Solana mainnet", {})],
            [("solana-devnet", {"color": LILAC}), ("  ← 本 demo 用的这条", {"color": INK3})],
        ],
        "凭证:AppId + AppSecret + Authorization ID 与私钥。Solana devnet 只有这条 provider 支持。",
    ),
    (
        "稳定币与 faucet",
        "USDC · 6 decimals",
        [
            [("devnet mint 4zMMC9srt5Ri5X14GAgXhaHii3GnPAEERYPJgZJDncDU", {})],
            [("faucet.circle.com", {})],
        ],
        "收款地址必须已有该代币账户,否则结算永远失败(转账不会替你创建)。",
    ),
]

REPOS: list[tuple[str, list[tuple[str, str]]]] = [
    (
        "参考实现",
        [
            (
                "aws-samples/sample-agentcore-cloudfront-x402-payments",
                "AWS 官方 AgentCore × x402 sample;付款环照它实现",
            ),
            (
                "krystiangw/agenticpay",
                "x402 付费 MCP server;本 demo 的 merchant 由它改写成 Lambda",
            ),
            (
                "coinbase/x402",
                "协议本体 + facilitator + @x402/core · @x402/express · @x402/svm",
            ),
        ],
    ),
    (
        "探路与教程",
        [
            ("civicteam/x402-mcp", "托管的 x402 MCP demo(Base Sepolia),用来抓真实 402 报文"),
            (
                "awslabs/agentcore-samples",
                "AgentCore 官方示例仓库(原 amazon-bedrock-agentcore-samples 重定向到此)",
            ),
        ],
    ),
    (
        "规范与文档",
        [
            ("modelcontextprotocol.io", "MCP 规范:JSON-RPC 2.0 报文 + streamable HTTP 传输"),
            (
                "x402.org · docs.aws.amazon.com/bedrock-agentcore",
                "x402 协议站与 facilitator;AgentCore Payments / Gateway 文档",
            ),
        ],
    ),
]


def slide_one(d: Deck) -> None:
    d.slide()
    d.head("MERCHANT SIDE  ·  WHAT YOU MUST SHIP", "一个付费 MCP,技术上要提供什么")
    d.text(
        e(0.72),
        e(1.55),
        e(4.6),
        e(0.7),
        [
            [
                ("协议、传输、工具形状都还是 MCP —— 多出来的只有", {"size": 10.5, "color": INK2}),
                ("调用之前先付钱", {"size": 10.5, "color": INK, "bold": True}),
                (
                    "。落到实现上是这六项;紫色编号的三项是普通 MCP 没有的。",
                    {"size": 10.5, "color": INK2},
                ),
            ]
        ],
        space=1.35,
    )
    y, rowh, colw = e(2.35), e(0.76), e(5.05)
    for num, title, desc, mark in REQS:
        d.box(e(0.72), y, colw, Emu(1), fill=LINE)
        if mark:
            d.box(e(0.72), y + e(0.08), e(0.34), e(0.2), fill=PURPLE)
            d.text(
                e(0.72),
                y + e(0.09),
                e(0.34),
                e(0.2),
                [num],
                size=9,
                color=WHITE,
                font=MONO,
                align=PP_ALIGN.CENTER,
            )
        else:
            d.text(e(0.72), y + e(0.09), e(0.34), e(0.2), [num], size=9, color=INK3, font=MONO)
        x2 = e(0.72) + e(0.5)
        d.text(x2, y + e(0.06), colw - e(0.5), e(0.26), [[(title, {"size": 13, "font": DISP})]])
        d.text(
            x2,
            y + e(0.32),
            colw - e(0.5),
            e(0.4),
            [[(desc, {"size": 9, "color": INK2})]],
            space=1.3,
        )
        y = y + rowh
    d.box(e(0.72), y, colw, Emu(1), fill=LINE)

    cx, cw, cy = e(6.15), d.w - e(6.15) - e(0.72), e(1.5)
    d.box(cx, cy, cw, e(5.05), fill=CODEBG, line=PURPLE_DIM)
    cy = cy + e(0.24)
    line_h = 0.185
    for badge, cap, lines, highlights in SNIPPETS:
        d.box(cx + e(0.26), cy, e(0.3), e(0.18), fill=PURPLE)
        d.text(
            cx + e(0.26),
            cy + e(0.01),
            e(0.3),
            e(0.18),
            [badge],
            size=8,
            color=WHITE,
            font=MONO,
            align=PP_ALIGN.CENTER,
        )
        d.text(
            cx + e(0.64),
            cy,
            cw,
            e(0.18),
            [cap],
            size=8.5,
            color=RGBColor(0xB9, 0xAE, 0xDC),
            font=MONO,
        )
        cy = cy + e(0.26)
        for i, code in enumerate(lines):
            ly = cy + e(line_h * i)
            if i in highlights:
                d.box(cx + e(0.2), ly - e(0.012), cw - e(0.4), e(line_h), fill=HLBG)
                d.box(cx + e(0.2), ly - e(0.012), e(0.022), e(line_h), fill=PURPLE)
            d.text(cx + e(0.26), ly, cw - e(0.4), e(line_h), [colorize(code)], space=1.0)
        cy = cy + e(line_h * len(lines)) + e(0.2)
        d.box(cx + e(0.26), cy, cw - e(0.52), Emu(1), fill=RGBColor(0x3A, 0x2F, 0x55))
        cy = cy + e(0.2)
    d.footer("FACILITATOR: X402.ORG 或自建 · 测试网 BASE SEPOLIA / SOLANA DEVNET", "01 / 03")


def slide_two(d: Deck) -> None:
    d.slide()
    d.head("PAID MCP vs ORDINARY MCP", "四个假设被“先付钱”打破")
    c1, c2, c3 = e(0.72), e(0.72) + e(1.6), e(0.72) + e(6.3)
    for x, label, color in (
        (c1, "维度", INK3),
        (c2, "普通 MCP 的假设", INK3),
        (c3, "→  付费 MCP 换成了什么", LILAC),
    ):
        d.text(x, e(1.62), e(4.6), e(0.22), [label], size=9, color=color, font=MONO)
    y, rh = e(1.92), e(1.24)
    for num, dim, was, now, tag in ASSUMPTIONS:
        d.box(c1, y, d.w - e(1.44), Emu(1), fill=LINE)
        d.text(
            c1, y + e(0.1), e(1.2), e(0.6), [[(num, {"size": 38, "color": GHOST, "font": DISP})]]
        )
        d.text(
            c1 + e(0.42),
            y + e(0.28),
            e(1.2),
            e(0.3),
            [[(dim, {"size": 14.5, "font": DISP})]],
        )
        d.text(c2, y + e(0.16), e(4.5), e(0.9), [[(was, {"size": 11, "color": INK3})]], space=1.35)
        d.box(c3 - e(0.16), y + e(0.1), d.w - c3 - e(0.56), e(0.92), fill=NAVY)
        d.box(c3 - e(0.16), y + e(0.1), e(0.03), e(0.92), fill=PURPLE)
        d.text(c3, y + e(0.16), e(5.6), e(0.9), [[(now, {"size": 11.5})]], space=1.35)
        d.text(c3, y + e(0.9), e(5.6), e(0.22), [tag], size=9, color=MINT, font=MONO)
        y = y + rh
    d.box(c1, y, d.w - e(1.44), Emu(1), fill=LINE)
    d.footer("协议没变 · 传输没变 · 变的是“调用之前先付钱”", "02 / 03")


def slide_three(d: Deck) -> None:
    d.slide()
    d.head("WALLETS · CHAINS · RESOURCES", "用什么钱包,跑在哪条链")

    left, lw = e(0.72), e(5.6)
    y = e(1.66)
    d.text(
        left, y, lw, e(0.24), [[("支持的钱包与对应链", {"size": 10, "color": LILAC, "font": MONO})]]
    )
    y = y + e(0.3)
    d.box(left, y, lw, Emu(1), fill=LINE)
    y = y + e(0.18)
    for name, meta, chain_lines, note in WALLETS:
        d.text(left, y, lw, e(0.3), [[(name, {"size": 15, "font": DISP})]])
        y = y + e(0.3)
        d.text(left, y, lw, e(0.22), [[(meta, {"size": 9, "color": INK3, "font": MONO})]])
        y = y + e(0.26)
        # chain ids as plain mono text: boxes around measured CJK text overflow at these sizes
        styled = [
            [(t, {"size": 10, "font": MONO, "color": st.get("color", INK2)}) for t, st in ln]
            for ln in chain_lines
        ]
        d.text(left, y, lw, e(0.2 * len(styled) + 0.1), styled, space=1.45)
        y = y + e(0.2 * len(styled) + 0.08)
        d.text(left, y, lw, e(0.42), [[(note, {"size": 9.5, "color": INK2})]], space=1.35)
        y = y + e(0.4)
        d.box(left, y, lw, Emu(1), fill=LINE)
        y = y + e(0.2)

    right = e(0.72) + e(6.1)
    rw = d.w - right - e(0.72)
    y = e(1.5)
    d.text(
        right, y, rw, e(0.24), [[("参考 REPO 与文档", {"size": 10, "color": LILAC, "font": MONO})]]
    )
    y = y + e(0.3)
    d.box(right, y, rw, Emu(1), fill=LINE)
    y = y + e(0.12)
    for group, items in REPOS:
        d.text(right, y, rw, e(0.2), [[(group, {"size": 8.5, "color": INK3, "font": MONO})]])
        y = y + e(0.24)
        for name, desc in items:
            d.box(right, y, rw, Emu(1), fill=LINE)
            y = y + e(0.09)
            d.text(right, y, rw, e(0.24), [[(name, {"size": 10.5, "font": MONO})]])
            y = y + e(0.23)
            d.text(right, y, rw, e(0.34), [[(desc, {"size": 9.5, "color": INK2})]], space=1.3)
            y = y + e(0.29)
        y = y + e(0.03)
    d.footer("COINBASE CDP 与 STRIPE(PRIVY) · X402 V1/V2 · MPP", "03 / 03")


def main() -> None:
    deck = Deck()
    slide_one(deck)
    slide_two(deck)
    slide_three(deck)
    deck.prs.save(OUT)
    print(f"wrote {OUT} ({len(deck.prs.slides)} slides)")


if __name__ == "__main__":
    main()
